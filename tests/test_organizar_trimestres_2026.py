from __future__ import annotations

import os
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

import organizar_trimestres_2026 as mod


class OrganizarTrimestres2026Tests(unittest.TestCase):
    def _excel(self, path: Path, rows_turnados=(), rows_internos=()):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Turnados recibidos"
        ws.append(["1711", "Fecha de creación", "Ruta"])
        for row in rows_turnados:
            ws.append(list(row))
        wi = wb.create_sheet("Internos")
        wi.append(["1711", "Fecha de creación", "Ruta"])
        for row in rows_internos:
            wi.append(list(row))
        wb.save(path)
        wb.close()

    def test_mapeo_personalizado_fechas_y_limites(self):
        self.assertEqual(mod.bucket_para_fecha(date(2026, 10, 1)), "2026Q3")
        self.assertEqual(mod.bucket_para_fecha(date(2026, 12, 15)), "2026Q3")
        self.assertIsNone(mod.bucket_para_fecha(date(2026, 12, 16)))
        self.assertIsNone(mod.bucket_para_fecha(date(2026, 1, 1)))
        self.assertIsNone(mod.bucket_para_fecha(date(2026, 3, 31)))
        self.assertEqual(mod.bucket_para_fecha(date(2027, 1, 1)), "2026Q4")
        self.assertEqual(mod.bucket_para_fecha(date(2027, 3, 31)), "2026Q4")
        self.assertIsNone(mod.bucket_para_fecha(date(2027, 4, 1)))

    def test_copia_ruta_excel_preserva_subcarpetas_y_no_publica_json(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                ("CRT26-000160", "05/01/2027", "521331_linktech_enlaces_de_tecnologia_s_a_de_c_v\\01 EN\\VE"),
            ])
            src = root / "descargas" / "CRT26-000160"
            (src / "anexos").mkdir(parents=True)
            (src / "solicitud.pdf").write_bytes(b"NUEVO")
            (src / "anexos" / "r001.xlsx").write_bytes(b"X")
            (src / "metadata_satys.json").write_text("{}", encoding="utf-8")
            output = root / "output"
            destino = output / "2026Q4" / "521331_linktech_enlaces_de_tecnologia_s_a_de_c_v" / "01 EN" / "VE"
            destino.mkdir(parents=True)
            (destino / "solicitud.pdf").write_bytes(b"VIEJO")
            (destino / "historico.txt").write_bytes(b"H")
            shared = root / "shared"
            old = os.environ.get("SATYS_REQUIRE_SHARED_MOUNT")
            os.environ["SATYS_REQUIRE_SHARED_MOUNT"] = "0"
            try:
                payload = mod.organizar(
                    excel_path=excel,
                    descargas_base=root / "descargas",
                    output_base=output,
                    shared_root=shared,
                    logs_dir=root / "logs",
                )
            finally:
                if old is None:
                    os.environ.pop("SATYS_REQUIRE_SHARED_MOUNT", None)
                else:
                    os.environ["SATYS_REQUIRE_SHARED_MOUNT"] = old
            self.assertEqual(payload["seleccionados"]["2026Q4"], 1)
            normal = output / "521331_linktech_enlaces_de_tecnologia_s_a_de_c_v" / "01 EN" / "VE"
            self.assertEqual((normal / "solicitud.pdf").read_bytes(), b"NUEVO")
            self.assertEqual((destino / "solicitud.pdf").read_bytes(), b"NUEVO")
            self.assertEqual((destino / "historico.txt").read_bytes(), b"H")
            self.assertEqual((destino / "anexos" / "r001.xlsx").read_bytes(), b"X")
            self.assertFalse(any(destino.rglob("*.json")))
            self.assertFalse((destino / "solicitud_1.pdf").exists())
            remoto = shared / "output" / "2026Q4" / "521331_linktech_enlaces_de_tecnologia_s_a_de_c_v" / "01 EN" / "VE"
            self.assertEqual((remoto / "solicitud.pdf").read_bytes(), b"NUEVO")
            self.assertEqual((remoto / "historico.txt").read_bytes(), b"H")
            remoto_normal = shared / "output" / "521331_linktech_enlaces_de_tecnologia_s_a_de_c_v" / "01 EN" / "VE"
            self.assertEqual((remoto_normal / "solicitud.pdf").read_bytes(), b"NUEVO")

            wb = openpyxl.load_workbook(excel, read_only=True, data_only=False)
            try:
                self.assertEqual(
                    wb["Turnados recibidos"]["C2"].value,
                    r"2026Q4\521331_linktech_enlaces_de_tecnologia_s_a_de_c_v\01 EN\VE",
                )
            finally:
                wb.close()
            wb = openpyxl.load_workbook(shared / "TrámitesCRT.xlsx", read_only=True, data_only=False)
            try:
                self.assertEqual(
                    wb["Turnados recibidos"]["C2"].value,
                    r"2026Q4\521331_linktech_enlaces_de_tecnologia_s_a_de_c_v\01 EN\VE",
                )
            finally:
                wb.close()
            self.assertEqual(payload["rutas_excel_actualizadas"], 1)
            self.assertTrue(payload["excel_depi_actualizado"])

    def test_octubre_a_diciembre_va_a_2026q3(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                ("CRT26-099001", "01-10-2026 00:00:00", "519564_starlink_satellite_systems_mexico_s_de_r_l_de_c_v\\01 EN\\VE"),
                ("CRT26-099002", "15/12/2026", "519564_starlink_satellite_systems_mexico_s_de_r_l_de_c_v\\01 EN\\VE"),
            ])
            for reg, contenido in (("CRT26-099001", b"A"), ("CRT26-099002", b"B")):
                d = root / "descargas" / reg
                d.mkdir(parents=True)
                (d / f"{reg}.pdf").write_bytes(contenido)
            payload = mod.organizar(
                excel_path=excel,
                descargas_base=root / "descargas",
                output_base=root / "output",
                shared_root=root / "shared",
                logs_dir=root / "logs",
                sincronizar_depi=False,
            )
            self.assertEqual(payload["seleccionados"]["2026Q3"], 2)
            dest = root / "output" / "2026Q3" / "519564_starlink_satellite_systems_mexico_s_de_r_l_de_c_v" / "01 EN" / "VE"
            self.assertTrue((dest / "CRT26-099001.pdf").exists())
            self.assertTrue((dest / "CRT26-099002.pdf").exists())
            wb = openpyxl.load_workbook(excel, read_only=True, data_only=False)
            try:
                self.assertTrue(str(wb["Turnados recibidos"]["C2"].value).startswith("2026Q3\\"))
                self.assertTrue(str(wb["Turnados recibidos"]["C3"].value).startswith("2026Q3\\"))
            finally:
                wb.close()

    def test_internos_encuentra_fuentes_exactas_en_varias_bandejas(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_internos=[
                ("157600", "23/01/2027", "_sin_operador\\internos__Ultimos_Movimientos__157600"),
            ])
            a = root / "descargas" / "internos" / "atendidos" / "157600"
            b = root / "descargas" / "internos" / "fuera_de_tiempo" / "157600"
            a.mkdir(parents=True); b.mkdir(parents=True)
            (a / "uno.pdf").write_bytes(b"1")
            (b / "dos.pdf").write_bytes(b"2")
            payload = mod.organizar(
                excel_path=excel,
                descargas_base=root / "descargas",
                output_base=root / "output",
                shared_root=root / "shared",
                logs_dir=root / "logs",
                sincronizar_depi=False,
            )
            self.assertEqual(payload["expedientes_con_fuente"], 1)
            dest = root / "output" / "2026Q4" / "_sin_operador" / "internos__Ultimos_Movimientos__157600"
            self.assertTrue((dest / "uno.pdf").exists())
            self.assertTrue((dest / "dos.pdf").exists())

    def test_readonly_no_usa_acceso_aleatorio_ws_cell(self):
        # Regresión 2026-09-08: ReadOnlyWorksheet.cell() vuelve a parsear el XML
        # desde el inicio. Llamarlo por cada fila hacía que el dry-run pareciera
        # congelado durante horas. La selección debe ser completamente streaming.
        from unittest.mock import patch

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                (f"CRT26-{i:06d}", "05/01/2027", "operador\\01 EN\\VE")
                for i in range(1, 50)
            ])
            (root / "descargas").mkdir()
            with patch(
                "openpyxl.worksheet._read_only.ReadOnlyWorksheet.cell",
                side_effect=AssertionError("no usar ws.cell en ReadOnlyWorksheet"),
            ):
                payload = mod.organizar(
                    excel_path=excel,
                    descargas_base=root / "descargas",
                    output_base=root / "output",
                    shared_root=root / "shared",
                    logs_dir=root / "logs",
                    dry_run=True,
                    sincronizar_depi=False,
                )
            self.assertEqual(payload["total_seleccionados"], 49)

    def test_dry_run_no_escribe(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                ("CRT26-000160", "05/01/2027", "operador\\01 EN\\VE"),
            ])
            src = root / "descargas" / "CRT26-000160"
            src.mkdir(parents=True)
            (src / "a.pdf").write_bytes(b"A")
            payload = mod.organizar(
                excel_path=excel,
                descargas_base=root / "descargas",
                output_base=root / "output",
                shared_root=root / "shared",
                logs_dir=root / "logs",
                dry_run=True,
            )
            self.assertEqual(payload["total_seleccionados"], 1)
            self.assertFalse((root / "output" / "2026Q4").exists())
            self.assertFalse((root / "shared" / "output").exists())

    def test_ruta_absoluta_antigua_se_vuelve_relativa_a_bucket(self):
        self.assertEqual(
            mod._ruta_relativa_segura(r"/app/output/519564_starlink_satellite_systems_mexico_s_de_r_l_de_c_v\01 EN\VE"),
            Path("519564_starlink_satellite_systems_mexico_s_de_r_l_de_c_v") / "01 EN" / "VE",
        )

    def test_no_duplica_prefijo_de_bucket(self):
        self.assertEqual(
            mod._ruta_relativa_segura(r"2026Q4\operador\01 EN\VE"),
            Path("operador") / "01 EN" / "VE",
        )

    def test_rerun_no_duplica_prefijo_y_mantiene_dos_destinos(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                ("CRT27-000001", "05/01/2027", r"2026Q4\operador_demo\01 EN\VE"),
            ])
            src = root / "descargas" / "CRT27-000001"
            src.mkdir(parents=True)
            (src / "a.pdf").write_bytes(b"A")
            payload = mod.organizar(
                excel_path=excel,
                descargas_base=root / "descargas",
                output_base=root / "output",
                shared_root=root / "shared",
                logs_dir=root / "logs",
                sincronizar_depi=False,
            )
            self.assertEqual(payload["rutas_excel_ya_correctas"], 1)
            self.assertTrue((root / "output" / "operador_demo" / "01 EN" / "VE" / "a.pdf").exists())
            self.assertTrue((root / "output" / "2026Q4" / "operador_demo" / "01 EN" / "VE" / "a.pdf").exists())
            self.assertFalse((root / "output" / "2026Q4" / "2026Q4").exists())

    def test_fallo_depi_no_confirma_ruta_excel(self):
        from unittest.mock import patch

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                ("CRT27-000002", "10/02/2027", r"operador_demo\01 EN\VE"),
            ])
            src = root / "descargas" / "CRT27-000002"
            src.mkdir(parents=True)
            (src / "a.pdf").write_bytes(b"A")
            with patch.object(mod, "validar_destino_compartido", return_value="DEPI no montado"):
                payload = mod.organizar(
                    excel_path=excel,
                    descargas_base=root / "descargas",
                    output_base=root / "output",
                    shared_root=root / "shared",
                    logs_dir=root / "logs",
                    sincronizar_depi=True,
                )
            wb = openpyxl.load_workbook(excel, read_only=True, data_only=False)
            try:
                self.assertEqual(wb["Turnados recibidos"]["C2"].value, r"operador_demo\01 EN\VE")
            finally:
                wb.close()
            self.assertEqual(payload["rutas_excel_actualizadas"], 0)
            self.assertEqual(payload["rutas_excel_bloqueadas_depi"], 1)
            self.assertTrue(payload["errores"])
            self.assertTrue((root / "output" / "operador_demo" / "01 EN" / "VE" / "a.pdf").exists())
            self.assertTrue((root / "output" / "2026Q4" / "operador_demo" / "01 EN" / "VE" / "a.pdf").exists())

    def test_fuera_de_ventana_enero_2026_no_se_mueve_ni_cambia_excel(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            excel = root / "TrámitesCRT.xlsx"
            self._excel(excel, rows_turnados=[
                ("CRT26-000160", "05/01/2026", r"operador\01 EN\VE"),
            ])
            src = root / "descargas" / "CRT26-000160"
            src.mkdir(parents=True)
            (src / "a.pdf").write_bytes(b"A")
            payload = mod.organizar(
                excel_path=excel,
                descargas_base=root / "descargas",
                output_base=root / "output",
                shared_root=root / "shared",
                logs_dir=root / "logs",
                sincronizar_depi=False,
            )
            self.assertEqual(payload["total_seleccionados"], 0)
            self.assertFalse((root / "output" / "2026Q4").exists())
            wb = openpyxl.load_workbook(excel, read_only=True, data_only=False)
            try:
                self.assertEqual(wb["Turnados recibidos"]["C2"].value, r"operador\01 EN\VE")
            finally:
                wb.close()

    def test_corrida_diaria_coloca_etapa_antes_del_correo(self):
        root = Path(mod.__file__).resolve().parent
        fuente = (root / "automatizar_registros_diario.py").read_text(encoding="utf-8")
        pos_etapa = fuente.index("# Organización final solicitada")
        pos_correo = fuente.index("# Los dos main_procesar.py siempre reciben --sin-email", pos_etapa)
        self.assertLess(pos_etapa, pos_correo)

    def test_wrapper_tiene_comando_independiente(self):
        root = Path(mod.__file__).resolve().parent
        fuente = (root / "scripts" / "podman_satys.sh").read_text(encoding="utf-8")
        self.assertIn("trimestres-2026)", fuente)
        self.assertIn("python organizar_trimestres_2026.py", fuente)


if __name__ == "__main__":
    unittest.main()
