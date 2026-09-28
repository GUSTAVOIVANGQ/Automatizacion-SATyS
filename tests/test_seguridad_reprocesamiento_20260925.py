from __future__ import annotations

import json
import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

import main_procesar as main
import Parte4_excel as p4
import reconciliar_metadata_global as global_recon


class FiltrosDiariosSegurosTests(unittest.TestCase):
    def test_turnados_solo_procesa_registros_autorizados(self):
        candidatos = [
            (Path('/tmp/a'), 'a', 'CRT26-000001'),
            (Path('/tmp/b'), 'b', 'CRT26-000002'),
            (Path('/tmp/c'), 'c', 'CRT25-000003'),
        ]
        filtrados, encontrados = main.filtrar_descargas_por_registros(
            candidatos, {'CRT26-000002'}
        )
        self.assertEqual([x[2] for x in filtrados], ['CRT26-000002'])
        self.assertEqual(encontrados, {'CRT26-000002'})

    def test_internos_solo_procesa_pares_bandeja_folio_objetivo(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            cand = []
            for bandeja, folio in [
                ('Atendidos', '100001'),
                ('En proceso', '100001'),
                ('Atendidos', '100002'),
            ]:
                carpeta = root / bandeja / folio
                carpeta.mkdir(parents=True)
                (carpeta / 'metadata_satys.json').write_text(json.dumps({
                    'bandeja_internos': bandeja,
                    'folio_tabla_internos': folio,
                    'folio': folio,
                }), encoding='utf-8')
                cand.append((carpeta, f'internos__{bandeja}__{folio}', folio))

            filtrados, encontrados = main.filtrar_descargas_internos_por_objetivos(
                cand, {(main.slug_bandeja_internos('Atendidos'), '100001')}
            )
            self.assertEqual(len(filtrados), 1)
            self.assertEqual(filtrados[0][0].parent.name, 'Atendidos')
            self.assertEqual(encontrados, {(main.slug_bandeja_internos('Atendidos'), '100001')})


class ExcelConservadorTests(unittest.TestCase):
    def _excel(self, path: Path, ruta: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Turnados recibidos'
        ws.append([
            '1711', 'Memo/Volante', 'Solicitante Promovente', 'Representante Legal',
            'Asunto', 'Tipo Trámite', 'Fecha de creación', 'FECHA LÍMITE', 'Ruta',
            'R001', 'R026', 'NOTAS_VICTOR',
        ])
        ws.append([
            'CRT26-000001', '123', 'OPERADOR VALIDADO', 'REP VALIDADO',
            'ASUNTO VALIDADO', 'TIPO VALIDADO', '01/01/2026', '02/01/2026', ruta,
            1, None, '',
        ])
        wb.save(path)
        wb.close()

    def test_rpc_caido_no_degrada_ruta_ni_campos_existentes(self):
        with tempfile.TemporaryDirectory() as td:
            excel = Path(td) / 'TrámitesCRT.xlsx'
            self._excel(excel, r'500_operador_validado\01 EN\VE')
            ok = p4.actualizar_excel(
                folio='123', registro='CRT26-000001',
                nombre_operador='NOMBRE PARCIAL', representante_legal='REP PARCIAL',
                formatos={'R026': True}, rpc_resultado={'ok': False},
                fecha_sello='09/09/2026', excel_path=excel,
                asunto='ASUNTO NUEVO', tipo_tramite='TIPO NUEVO',
                fecha_limite='10/10/2026', ruta_salida=r'_sin_operador\CRT26-000001',
            )
            self.assertTrue(ok)
            wb = openpyxl.load_workbook(excel, data_only=False)
            ws = wb['Turnados recibidos']
            self.assertEqual(ws['C2'].value, 'OPERADOR VALIDADO')
            self.assertEqual(ws['D2'].value, 'REP VALIDADO')
            self.assertEqual(ws['E2'].value, 'ASUNTO VALIDADO')
            self.assertEqual(ws['F2'].value, 'TIPO VALIDADO')
            self.assertEqual(ws['G2'].value, '01/01/2026')
            self.assertEqual(ws['H2'].value, '02/01/2026')
            self.assertEqual(ws['I2'].value, r'500_operador_validado\01 EN\VE')
            self.assertEqual(ws['J2'].value, 1)   # R001 no se limpia
            self.assertEqual(ws['K2'].value, 1)   # R026 sí se agrega
            wb.close()

    def test_ruta_pendiente_si_puede_promoverse_a_canonica(self):
        with tempfile.TemporaryDirectory() as td:
            excel = Path(td) / 'TrámitesCRT.xlsx'
            self._excel(excel, r'_sin_operador\CRT26-000001')
            ok = p4.actualizar_excel(
                folio='123', registro='CRT26-000001',
                rpc_resultado={'ok': True, 'ruta': r'500_operador\01 EN\VE'},
                excel_path=excel, ruta_salida=r'500_operador\01 EN\VE',
            )
            self.assertTrue(ok)
            wb = openpyxl.load_workbook(excel, data_only=False)
            self.assertEqual(wb['Turnados recibidos']['I2'].value, r'500_operador\01 EN\VE')
            wb.close()


class CatalogoYGlobalSeguroTests(unittest.TestCase):
    def test_catalogo_global_ignora_xlsx_mas_nuevo_corrupto(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            viejo = base / '03_concesiones_permisos_autorizaciones_20260901.xlsx'
            nuevo = base / '03_concesiones_permisos_autorizaciones_20260925.xlsx'
            viejo.write_bytes(b'viejo')
            nuevo.write_bytes(b'corrupto')
            now = time.time()
            os.utime(viejo, (now - 100, now - 100))
            os.utime(nuevo, (now, now))

            def cargar(path, *_args, **_kwargs):
                if Path(path).name == nuevo.name:
                    raise ValueError('File is not a zip file')
                return [{'idBp': '500', 'nombre_completo': 'OPERADOR'}]

            with patch.object(global_recon.bc, 'cargar_catalogo_desde_excel', side_effect=cargar), \
                 patch.object(global_recon.bc, 'preparar_catalogo_para_matching', side_effect=lambda x: x):
                indice, elegido = global_recon.cargar_indice_rpc(base)
            self.assertEqual(elegido, viejo)
            self.assertIn('500', indice)

    def test_global_no_consulta_rpc_para_ruta_canonica_existente(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            carpeta = root / 'descargas' / 'CRT26-000001'
            carpeta.mkdir(parents=True)
            (carpeta / 'metadata_satys.json').write_text(json.dumps({
                'registro': 'CRT26-000001', 'folio': '1',
                'id_solicitante': '500', 'nombre_operador': 'OPERADOR',
            }), encoding='utf-8')
            with patch.object(global_recon.bc, 'resolver_operador_seguro', side_effect=AssertionError('no debe consultarse')):
                resultados, _ = global_recon.construir_resultados(
                    root / 'descargas', root / 'output', {},
                    rutas_existentes={'CRT26-000001': r'500_operador\01 EN\VE'},
                    reorganizar_output=False,
                )
            self.assertEqual(resultados[0]['rpc_resultado']['metodo'], 'ruta_excel_preservada')
            self.assertTrue(resultados[0]['rpc_ok'])


if __name__ == '__main__':
    unittest.main()
