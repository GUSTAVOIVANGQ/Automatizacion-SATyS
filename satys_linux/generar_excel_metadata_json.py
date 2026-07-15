#!/usr/bin/env python3
r"""
generar_excel_metadata_json.py
──────────────────────────────
Genera output/Folios_Datos_Completos.xlsx con todos los campos de:
  - metadata_satys.json
  - metadata_tramite_nuevo.json
por cada carpeta/número de registro procesado.

También agrega las columnas:
  - output: ruta relativa donde quedaron organizados los archivos
  - descargas: ruta relativa de la carpeta fuente en descargas/
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise RuntimeError("Falta openpyxl. Instala con: python -m pip install openpyxl") from exc

log = logging.getLogger("SATyS-ExcelMetadata")

JSON_NAMES = ("metadata_satys.json", "metadata_tramite_nuevo.json")


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception as exc:
        log.warning("No se pudo leer %s: %s", path, exc)
        return {}


def _rel_backslash(path_value: str | Path | None, project_root: Path | None = None) -> str:
    if not path_value:
        return ""
    p = Path(path_value)
    root = project_root or Path.cwd()
    try:
        rel = p.resolve().relative_to(root.resolve())
    except Exception:
        try:
            rel = p.relative_to(root)
        except Exception:
            rel = p
    text = str(rel).replace("/", "\\")
    if not text.startswith("\\"):
        text = "\\" + text
    return text


def _registro_from(meta_satys: dict[str, Any], meta_tn: dict[str, Any], fallback: str = "") -> str:
    for key in ("registro", "numero_registro", "1711"):
        val = meta_satys.get(key) or meta_tn.get(key)
        if val:
            return str(val).strip().upper()
    return str(fallback or "").strip().upper()


def _folio_from(meta_satys: dict[str, Any], meta_tn: dict[str, Any], fallback: str = "") -> str:
    for key in ("folio", "folio_opc", "memo_folio_opc"):
        val = meta_satys.get(key) or meta_tn.get(key)
        if val:
            return str(val).strip()
    return str(fallback or "").strip()


def _descargas_from_result(resultado: dict[str, Any], descargas_base: Path) -> Path:
    for key in ("descargas_dir", "carpeta", "carpeta_descarga"):
        val = resultado.get(key)
        if val:
            return Path(val)
    folio_id = resultado.get("folio_id") or resultado.get("registro") or resultado.get("folio")
    return descargas_base / str(folio_id or "")


def _output_from_result(resultado: dict[str, Any], output_base: Path) -> Path:
    for key in ("output_dir", "sin_operador_dir"):
        val = resultado.get(key)
        if val:
            return Path(val)
    rpc = resultado.get("rpc_resultado") or {}
    if isinstance(rpc, dict) and rpc.get("ok") and rpc.get("ruta"):
        return output_base / str(rpc["ruta"]).replace("\\", "/")
    folio_id = resultado.get("folio_id") or resultado.get("registro") or resultado.get("folio")
    return output_base / "_sin_operador" / str(folio_id or "")


def _scan_descargas(descargas_base: Path) -> list[dict[str, Any]]:
    resultados: list[dict[str, Any]] = []
    if not descargas_base.exists():
        return resultados
    for carpeta in sorted([p for p in descargas_base.rglob("*") if p.is_dir()], key=lambda p: str(p).upper()):
        if any((carpeta / name).exists() for name in JSON_NAMES):
            resultados.append({
                "folio": carpeta.name,
                "folio_id": carpeta.name,
                "descargas_dir": str(carpeta),
            })
    return resultados


def generar_excel_metadata_json(
    resultados: list[dict[str, Any]] | None = None,
    descargas_base: str | Path = "descargas",
    output_base: str | Path = "output",
    excel_salida: str | Path | None = None,
    project_root: str | Path | None = None,
) -> Path:
    """
    Genera el Excel consolidado de metadatos JSON.

    Usa 'resultados' de main_procesar.py para ubicar output/_sin_operador con
    precisión. Si no se pasan resultados, escanea descargas/.
    """
    descargas_base = Path(descargas_base)
    output_base = Path(output_base)
    project_root = Path(project_root) if project_root else Path.cwd()
    excel_salida = Path(excel_salida) if excel_salida else output_base / "Folios_Datos_Completos.xlsx"
    excel_salida.parent.mkdir(parents=True, exist_ok=True)

    resultados = list(resultados or [])
    if not resultados:
        resultados = _scan_descargas(descargas_base)

    filas: list[dict[str, Any]] = []
    satys_keys: set[str] = set()
    tramite_keys: set[str] = set()
    vistos_descargas: set[str] = set()

    for resultado in resultados:
        carpeta = _descargas_from_result(resultado, descargas_base)
        try:
            key_carpeta = str(carpeta.resolve())
        except Exception:
            key_carpeta = str(carpeta)
        if key_carpeta in vistos_descargas:
            continue
        vistos_descargas.add(key_carpeta)

        meta_satys = _read_json(carpeta / "metadata_satys.json")
        meta_tn = _read_json(carpeta / "metadata_tramite_nuevo.json")
        if not meta_satys and not meta_tn:
            continue

        satys_keys.update(str(k) for k in meta_satys.keys())
        tramite_keys.update(str(k) for k in meta_tn.keys())

        registro = _registro_from(meta_satys, meta_tn, resultado.get("registro") or resultado.get("folio_id") or carpeta.name)
        folio = _folio_from(meta_satys, meta_tn, resultado.get("folio") or carpeta.name)
        rpc = resultado.get("rpc_resultado") if isinstance(resultado.get("rpc_resultado"), dict) else {}
        output_dir = _output_from_result(resultado, output_base)

        filas.append({
            "registro": registro,
            "folio": folio,
            "id_solicitante": meta_satys.get("id_solicitante") or meta_tn.get("id_solicitante") or resultado.get("id_solicitante") or "",
            "nombre_operador": meta_satys.get("nombre_operador") or meta_tn.get("nombre_operador") or resultado.get("nombre_operador") or "",
            "representante_legal": meta_satys.get("representante_legal") or meta_tn.get("representante_legal") or resultado.get("representante_legal") or "",
            "rpc_ok": bool(resultado.get("rpc_ok")),
            "rpc_exactitud": 100 if resultado.get("rpc_ok") else 0,
            "rpc_metodo": rpc.get("metodo") or "id_exacto",
            "rpc_id_operador": rpc.get("idBp") or rpc.get("numero_rpc") or "",
            "output": _rel_backslash(output_dir, project_root),
            "descargas": _rel_backslash(carpeta, project_root),
            "_meta_satys": meta_satys,
            "_meta_tramite_nuevo": meta_tn,
        })

    satys_cols = [f"metadata_satys.{k}" for k in sorted(satys_keys, key=str.lower)]
    tramite_cols = [f"metadata_tramite_nuevo.{k}" for k in sorted(tramite_keys, key=str.lower)]
    base_cols = [
        "registro", "folio", "id_solicitante", "nombre_operador", "representante_legal",
        "rpc_ok", "rpc_exactitud", "rpc_metodo", "rpc_id_operador", "output", "descargas",
    ]
    headers = base_cols + satys_cols + tramite_cols

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos_Completos"
    ws.append(headers)

    for fila in filas:
        row = []
        for h in headers:
            if h.startswith("metadata_satys."):
                v = fila["_meta_satys"].get(h.split(".", 1)[1], "")
            elif h.startswith("metadata_tramite_nuevo."):
                v = fila["_meta_tramite_nuevo"].get(h.split(".", 1)[1], "")
            else:
                v = fila.get(h, "")
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False)
            row.append(v)
        ws.append(row)

    # Estilo simple y legible
    header_fill = PatternFill("solid", fgColor="156E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2E3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if ws.max_row >= 2 and ws.max_column >= 1:
        table_ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName="TablaDatosCompletos", ref=table_ref)
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        ws.add_table(table)

    # Anchos acotados
    for col_idx, h in enumerate(headers, 1):
        letter = get_column_letter(col_idx)
        if h in {"output", "descargas"} or h.startswith("metadata_"):
            ws.column_dimensions[letter].width = 34
        elif h in {"nombre_operador", "representante_legal"}:
            ws.column_dimensions[letter].width = 32
        else:
            ws.column_dimensions[letter].width = min(max(len(h) + 2, 12), 24)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Hoja resumen
    ws2 = wb.create_sheet("Resumen")
    total = len(filas)
    ok = sum(1 for f in filas if f.get("rpc_ok"))
    sin = total - ok
    ws2.append(["Métrica", "Valor"])
    ws2.append(["Total registros con JSON", total])
    ws2.append(["RPC 100% por ID", ok])
    ws2.append(["RPC 0% / sin operador", sin])
    ws2.append(["Archivo generado", str(excel_salida)])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 60

    wb.save(excel_salida)
    wb.close()
    log.info("Excel de metadatos JSON generado: %s (%d registros)", excel_salida, len(filas))
    return excel_salida


if __name__ == "__main__":
    generar_excel_metadata_json()
