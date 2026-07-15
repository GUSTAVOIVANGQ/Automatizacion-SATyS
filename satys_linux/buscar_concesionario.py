#!/usr/bin/env python3
r"""
buscar_concesionario.py
───────────────────────
Compatibilidad Linux para el cruce original SATyS/RPC.

Regla de negocio:
  - metadata_satys.json aporta id_solicitante.
  - El Excel oficial del RPC aporta la columna ID OPERADOR.
  - Si id_solicitante == ID OPERADOR: coincidencia 100%.
  - Si no existe/no coincide: coincidencia 0%; NO se resuelve por nombre/fuzzy.

Este módulo conserva los nombres de funciones que main_procesar.py esperaba en
la versión Windows: cargar_catalogo_desde_excel, preparar_catalogo_para_matching,
buscar_por_id_solicitante y buscar_coincidencias.
"""

from __future__ import annotations

import logging
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:
    raise RuntimeError("Falta openpyxl. Instala con: python -m pip install openpyxl") from exc

log = logging.getLogger("SATyS-BuscarConcesionario")


def _sin_acentos(valor: Any) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFD", texto)
    return "".join(c for c in texto if unicodedata.category(c) != "Mn")


def normalizar_header(valor: Any) -> str:
    """Normaliza encabezados para comparar sin acentos ni signos."""
    texto = _sin_acentos(valor).upper().strip()
    texto = re.sub(r"[^A-Z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def normalizar_id(valor: Any) -> str:
    """Normaliza IDs de SATyS/Excel: 518858, 518858.0, ' 518858 '."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    if not texto:
        return ""
    if re.fullmatch(r"\d+\.0+", texto):
        texto = texto.split(".", 1)[0]
    return re.sub(r"\s+", "", texto).upper()


def normalizar_nombre(valor: Any) -> str:
    texto = _sin_acentos(valor).lower()
    texto = re.sub(r"[^a-z0-9\s]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _buscar_hoja(wb, nombre_preferido: str | None):
    if nombre_preferido and nombre_preferido in wb.sheetnames:
        return wb[nombre_preferido]
    if nombre_preferido:
        objetivo = normalizar_header(nombre_preferido)
        for nombre in wb.sheetnames:
            if normalizar_header(nombre) == objetivo:
                return wb[nombre]
    return wb[wb.sheetnames[0]]


def _detectar_encabezados(ws, max_filas: int = 30) -> tuple[int, dict[str, int]]:
    """Detecta fila de encabezados; debe encontrar columna ID OPERADOR."""
    claves_id = {
        "ID OPERADOR", "ID DE OPERADOR", "IDOPERADOR", "ID BP", "IDBP", "ID_BP",
        "ID CONCESIONARIO", "IDCONCESIONARIO",
    }
    mejor: tuple[int, dict[str, int], int] | None = None

    for row_idx in range(1, min(ws.max_row, max_filas) + 1):
        headers: dict[str, int] = {}
        puntaje = 0
        for col_idx in range(1, ws.max_column + 1):
            h = normalizar_header(ws.cell(row=row_idx, column=col_idx).value)
            if not h:
                continue
            headers[h] = col_idx
            h_sin = h.replace(" ", "")
            if h in claves_id or h_sin in {x.replace(" ", "") for x in claves_id}:
                puntaje += 100
            if any(k in h for k in ("OPERADOR", "CONCESIONARIO", "RAZON SOCIAL", "NOMBRE", "TITULAR", "DENOMINACION")):
                puntaje += 2
            if h in {"ESTATUS", "ESTADO", "VIGENCIA", "SITUACION"}:
                puntaje += 1
        if puntaje and (mejor is None or puntaje > mejor[2]):
            mejor = (row_idx, headers, puntaje)

    if mejor is None or mejor[2] < 100:
        raise ValueError("No se encontró encabezado con columna 'ID OPERADOR' en el Excel RPC")
    return mejor[0], mejor[1]


def _columna_por_alias(headers: dict[str, int], aliases: list[str]) -> int | None:
    aliases_norm = [normalizar_header(a) for a in aliases]
    aliases_sin = {a.replace(" ", "") for a in aliases_norm}
    for alias in aliases_norm:
        if alias in headers:
            return headers[alias]
    for h, col in headers.items():
        if h.replace(" ", "") in aliases_sin:
            return col
    return None


def cargar_catalogo_desde_excel(
    excel_path: str | Path,
    hoja: str = "copeau",
    solo_vigentes: bool = False,
) -> list[dict[str, Any]]:
    """
    Lee el Excel oficial del RPC y devuelve operadores normalizados.

    Requiere columna ID OPERADOR. Para el nombre busca, en orden flexible,
    columnas como OPERADOR, CONCESIONARIO, RAZON SOCIAL, DENOMINACION, TITULAR.
    """
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"No existe el Excel RPC: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        ws = _buscar_hoja(wb, hoja)
        header_row, headers = _detectar_encabezados(ws)

        col_id = _columna_por_alias(headers, [
            "ID OPERADOR", "ID DE OPERADOR", "IDOPERADOR", "ID BP", "IDBP", "ID_BP",
            "ID CONCESIONARIO", "IDCONCESIONARIO",
        ])
        if not col_id:
            raise ValueError("No se encontró la columna 'ID OPERADOR' en el Excel RPC")

        col_nombre = _columna_por_alias(headers, [
            "OPERADOR", "NOMBRE OPERADOR", "NOMBRE DEL OPERADOR", "CONCESIONARIO",
            "NOMBRE CONCESIONARIO", "NOMBRE DEL CONCESIONARIO", "RAZON SOCIAL", "RAZÓN SOCIAL",
            "DENOMINACION", "DENOMINACIÓN", "TITULAR", "NOMBRE", "NOMBRE COMPLETO",
        ])
        if not col_nombre:
            for h, col in headers.items():
                if col != col_id and any(k in h for k in ("OPERADOR", "CONCESIONARIO", "RAZON", "DENOMINACION", "TITULAR", "NOMBRE")):
                    col_nombre = col
                    break
        if not col_nombre:
            raise ValueError("No se encontró columna de nombre de operador/concesionario en el Excel RPC")

        col_estatus = _columna_por_alias(headers, ["ESTATUS", "ESTADO", "SITUACION", "SITUACIÓN", "VIGENCIA"])

        catalogo: list[dict[str, Any]] = []
        ids_vistos: set[str] = set()
        for row in range(header_row + 1, ws.max_row + 1):
            id_bp = normalizar_id(ws.cell(row=row, column=col_id).value)
            nombre = str(ws.cell(row=row, column=col_nombre).value or "").strip()
            if not id_bp or not nombre:
                continue

            if solo_vigentes and col_estatus:
                estatus = normalizar_header(ws.cell(row=row, column=col_estatus).value)
                if estatus and "VIGENTE" not in estatus:
                    continue

            if id_bp in ids_vistos:
                continue
            ids_vistos.add(id_bp)

            catalogo.append({
                "idBp": id_bp,
                "ID OPERADOR": id_bp,
                "concesionario": nombre,
                "nombre_completo": nombre,
                "_fila_excel_rpc": row,
                "_hoja_excel_rpc": ws.title,
            })

        log.info("Catálogo RPC cargado desde %s/%s: %d operadores", excel_path.name, ws.title, len(catalogo))
        return catalogo
    finally:
        wb.close()


def preparar_catalogo_para_matching(catalogo: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """
    Compatibilidad con main_procesar.py: agrega norm/compact para identificar
    que el catálogo viene de Excel RPC, aunque el cruce operativo sea por ID.
    """
    preparado: list[dict[str, Any]] = []
    for item in catalogo or []:
        id_bp = normalizar_id(item.get("idBp") or item.get("ID OPERADOR"))
        nombre = str(item.get("nombre_completo") or item.get("concesionario") or "").strip()
        if not id_bp or not nombre:
            continue
        norm = normalizar_nombre(nombre)
        nuevo = dict(item)
        nuevo.update({
            "idBp": id_bp,
            "ID OPERADOR": id_bp,
            "concesionario": nombre,
            "nombre_completo": nombre,
            "norm": norm,
            "compact": norm.replace(" ", ""),
        })
        preparado.append(nuevo)
    return preparado


def buscar_por_id_solicitante(id_solicitante: Any, catalogo: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Cruce exacto: id_solicitante de SATyS == ID OPERADOR del Excel RPC."""
    objetivo = normalizar_id(id_solicitante)
    if not objetivo:
        return None
    for item in catalogo or []:
        id_bp = normalizar_id(item.get("idBp") or item.get("ID OPERADOR"))
        if id_bp == objetivo:
            nombre = str(item.get("nombre_completo") or item.get("concesionario") or "").strip()
            return {
                "idBp": id_bp,
                "ID OPERADOR": id_bp,
                "concesionario": nombre,
                "nombre_completo": nombre,
                "score": 1.0,
            }
    return None


def buscar_coincidencias(nombre: str, catalogo: list[dict[str, Any]], top_n: int = 5):
    """
    Compatibilidad con código viejo. No se usa en producción porque el flujo
    corregido no permite fuzzy, pero se conserva para no romper utilerías.
    """
    q = normalizar_nombre(nombre)
    if not q:
        return []
    resultados = []
    for item in catalogo or []:
        norm = item.get("norm") or normalizar_nombre(item.get("nombre_completo") or item.get("concesionario"))
        score = SequenceMatcher(None, q, norm).ratio()
        resultados.append((score, item))
    resultados.sort(key=lambda x: x[0], reverse=True)
    return resultados[:top_n]
