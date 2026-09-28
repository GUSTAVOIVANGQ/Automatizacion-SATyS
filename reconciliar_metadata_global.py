#!/usr/bin/env python3
"""Reconciliación conservadora de Rutas desde metadatos locales.

La corrida diaria nunca reconstruye ni sobrescribe masivamente TrámitesCRT.xlsx.
Las Rutas canónicas existentes son inmutables; sólo se intenta promover filas
vacías/_sin_operador cuando existe evidencia exacta del RPC.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import buscar_concesionario as bc
import openpyxl
from Parte3_rpc import construir_ruta, construir_ruta_operadores
from Parte4_excel import (
    organizar_archivos,
    organizar_correo_exclusivo,
    normalizar_ruta_excel,
    ruta_es_canonica_estable,
    ruta_es_revision_manual,
)
from estado_descargas import iter_archivos_publicables_output
from configuracion_local import ruta_configurada
from generar_excel_metadata_json import generar_excel_metadata_json
from guardado_seguro import reemplazar_desde_temporal
from rutas_salida import (
    destino_sin_operador,
    es_folio_opc_correo,
    folio_opc_desde_metadata,
    ruta_relativa_sin_operador,
)

log = logging.getLogger("SATyS-ReconciliacionGlobal")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%H:%M:%S",
)

REGISTRO_RE = re.compile(r"\b[A-Z]{2,6}\d{2}-\d{3,}\b", re.IGNORECASE)
JSON_NAMES = ("metadata_satys.json", "metadata_tramite_nuevo.json")


def _es_metadata_internos(meta_satys: dict[str, Any], meta_tn: dict[str, Any]) -> bool:
    """Evita mezclar expedientes de Internos con la hoja Turnados recibidos."""
    for metadata in (meta_satys, meta_tn):
        if str(metadata.get("satys_flujo") or "").strip().lower() == "internos":
            return True
        if metadata.get("bandeja_internos") or metadata.get("folio_tabla_internos"):
            return True
    return False


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
        return value if isinstance(value, dict) else {}
    except Exception as exc:
        log.warning("No se pudo leer %s: %s", path, exc)
        return {}


def _registro(meta_satys: dict[str, Any], meta_tn: dict[str, Any], fallback: str) -> str:
    for key in ("registro", "numero_registro", "1711"):
        value = meta_satys.get(key) or meta_tn.get(key)
        text = str(value or "").strip().upper().replace(" ", "")
        match = REGISTRO_RE.search(text)
        if match:
            return match.group(0).upper()
    match = REGISTRO_RE.search(str(fallback or "").strip().upper())
    return match.group(0).upper() if match else ""


def _folio(meta_satys: dict[str, Any], meta_tn: dict[str, Any], fallback: str) -> str:
    directo = meta_satys.get("folio") or meta_tn.get("folio")
    if directo not in (None, ""):
        return str(directo).strip()
    folio_opc = folio_opc_desde_metadata(meta_satys, meta_tn)
    if folio_opc:
        numeros = re.sub(r"[^0-9]", "", folio_opc)
        if numeros:
            return numeros
    memo = meta_satys.get("memo_folio_opc") or meta_tn.get("memo_folio_opc")
    return str(memo or fallback or "").strip()


def _identificador_salida(carpeta: Path, descargas_base: Path, registro: str) -> str:
    """Replica la convención de main_procesar para carpetas normales/heredadas."""
    try:
        rel = carpeta.relative_to(descargas_base)
    except ValueError:
        return registro or carpeta.name
    if len(rel.parts) == 1:
        return rel.parts[0]
    return f"{rel.parts[0]}__{rel.parts[-1]}"


def _metadata_score(carpeta: Path, meta_satys: dict[str, Any], meta_tn: dict[str, Any]) -> tuple[int, float]:
    campos = sum(1 for value in list(meta_satys.values()) + list(meta_tn.values()) if value not in (None, ""))
    mtimes = [p.stat().st_mtime for p in (carpeta / JSON_NAMES[0], carpeta / JSON_NAMES[1]) if p.exists()]
    return campos, max(mtimes, default=0.0)


def descubrir_metadata(descargas_base: Path) -> tuple[list[dict[str, Any]], list[str]]:
    """Devuelve una carpeta canónica por Registro y reporta duplicados."""
    por_registro: dict[str, dict[str, Any]] = {}
    duplicados: list[str] = []
    if not descargas_base.exists():
        return [], []

    carpetas = sorted(
        {path.parent for name in JSON_NAMES for path in descargas_base.rglob(name)},
        key=lambda p: str(p).upper(),
    )
    for carpeta in carpetas:
        meta_satys = _read_json(carpeta / JSON_NAMES[0])
        meta_tn = _read_json(carpeta / JSON_NAMES[1])
        if not meta_satys and not meta_tn:
            continue
        if _es_metadata_internos(meta_satys, meta_tn):
            continue
        registro = _registro(meta_satys, meta_tn, carpeta.name)
        if not registro:
            log.warning("Se omite metadata sin Registro CRT válido: %s", carpeta)
            continue
        item = {
            "carpeta": carpeta,
            "meta_satys": meta_satys,
            "meta_tn": meta_tn,
            "registro": registro,
            "score": _metadata_score(carpeta, meta_satys, meta_tn),
        }
        anterior = por_registro.get(registro)
        if anterior is None or item["score"] > anterior["score"]:
            if anterior is not None:
                duplicados.append(registro)
            por_registro[registro] = item
        else:
            duplicados.append(registro)

    return [por_registro[k] for k in sorted(por_registro)], sorted(set(duplicados))


def _catalogos_rpc_por_recencia(base_rpc: Path) -> list[Path]:
    return sorted(
        base_rpc.glob("03_concesiones_permisos_autorizaciones_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )


def cargar_indice_rpc(base_rpc: Path) -> tuple[dict[str, dict[str, Any]], Path | None]:
    """Usa el catálogo RPC más reciente que sea realmente válido."""
    errores: list[str] = []
    for excel_rpc in _catalogos_rpc_por_recencia(base_rpc):
        try:
            catalogo = bc.cargar_catalogo_desde_excel(excel_rpc, "copeau", solo_vigentes=False)
            preparado = bc.preparar_catalogo_para_matching(catalogo)
            indice = {
                bc.normalizar_id(item.get("idBp")): item
                for item in preparado
                if bc.normalizar_id(item.get("idBp"))
            }
            if not indice:
                raise RuntimeError("catálogo vacío")
            if errores:
                log.warning(
                    "Se ignoraron %d Excel RPC más nuevos inválidos; se usa %s.",
                    len(errores), excel_rpc.name,
                )
            return indice, excel_rpc
        except Exception as exc:
            errores.append(f"{excel_rpc.name}: {exc}")
            log.warning("Excel RPC inválido; se ignora %s: %s", excel_rpc.name, exc)
    if errores:
        log.warning(
            "Ningún Excel RPC local es válido; sólo se intentará RPC en línea para pendientes: %s",
            " | ".join(errores[:5]),
        )
    else:
        log.warning("No existe Excel RPC local; sólo se intentará RPC en línea para pendientes.")
    return {}, None


def cargar_rutas_maestro(excel_path: Path) -> dict[str, str]:
    """Mapa Registro CRT -> Ruta actual, sin alterar el maestro."""
    if not excel_path.exists():
        return {}
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=False)
    try:
        if "Turnados recibidos" not in wb.sheetnames:
            return {}
        ws = wb["Turnados recibidos"]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        cols = {str(v or "").strip(): i for i, v in enumerate(header) if v not in (None, "")}
        col_reg = cols.get("1711")
        col_ruta = cols.get("Ruta")
        if col_reg is None or col_ruta is None:
            return {}
        out: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            reg = str(row[col_reg] or "").strip().upper() if col_reg < len(row) else ""
            if not REGISTRO_RE.fullmatch(reg):
                continue
            ruta = normalizar_ruta_excel(row[col_ruta] if col_ruta < len(row) else "")
            out[reg] = ruta
        return out
    finally:
        wb.close()


def aplicar_rutas_conservadoras(
    excel_path: Path,
    resultados: list[dict[str, Any]],
    *,
    crear_backup: bool = True,
) -> dict[str, Any]:
    """Aplica únicamente promociones seguras de Ruta, sin reconstruir filas.

    - vacío/_sin_operador -> ruta canónica: permitido;
    - ruta canónica existente -> jamás se sustituye automáticamente;
    - resultado no resuelto -> jamás degrada una ruta existente.
    """
    wb = openpyxl.load_workbook(excel_path)
    try:
        if "Turnados recibidos" not in wb.sheetnames:
            raise KeyError("No existe hoja Turnados recibidos")
        ws = wb["Turnados recibidos"]
        headers = {
            str(ws.cell(1, c).value or "").strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(1, c).value not in (None, "")
        }
        c_reg = headers.get("1711")
        c_ruta = headers.get("Ruta")
        if not c_reg or not c_ruta:
            raise ValueError("Faltan encabezados 1711/Ruta")
        filas = {}
        for r in range(2, ws.max_row + 1):
            reg = str(ws.cell(r, c_reg).value or "").strip().upper()
            if REGISTRO_RE.fullmatch(reg):
                filas.setdefault(reg, r)

        cambios = 0
        preservadas = 0
        pendientes = 0
        faltantes = 0
        for item in resultados:
            reg = str(item.get("registro") or "").strip().upper()
            r = filas.get(reg)
            if not r:
                faltantes += 1
                continue
            actual = normalizar_ruta_excel(ws.cell(r, c_ruta).value)
            nueva = normalizar_ruta_excel(
                (item.get("rpc_resultado") or {}).get("ruta")
                or item.get("ruta")
                or ""
            )
            if ruta_es_canonica_estable(actual):
                preservadas += 1
                continue
            if nueva and ruta_es_canonica_estable(nueva):
                ws.cell(r, c_ruta).value = nueva
                cambios += 1
            else:
                pendientes += 1

        backup = ""
        if cambios:
            if crear_backup:
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                backup_path = excel_path.with_name(
                    f"{excel_path.stem}_backup_pre_reconciliacion_segura_{stamp}{excel_path.suffix}"
                )
                shutil.copy2(excel_path, backup_path)
                backup = str(backup_path)
            temp = excel_path.with_name(f".{excel_path.name}.recon.tmp")
            wb.save(temp)
            reemplazar_desde_temporal(temp, excel_path)
        return {
            "routes_promoted": cambios,
            "canonical_preserved": preservadas,
            "still_pending": pendientes,
            "records_missing_in_master": faltantes,
            "backup": backup,
        }
    finally:
        wb.close()


def construir_resultados(
    descargas_base: Path,
    output_base: Path,
    indice_rpc: dict[str, dict[str, Any]],
    *,
    rutas_existentes: dict[str, str] | None = None,
    migrar_correos: bool = True,
    reorganizar_output: bool = True,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    items, duplicados = descubrir_metadata(descargas_base)
    resultados: list[dict[str, Any]] = []
    stats = {
        "metadata": len(items),
        "rpc_ok": 0,
        "sin_operador": 0,
        "sin_operador_correo": 0,
        "archivos_correo_copiados": 0,
        "archivos_organizados": 0,
        "metadata_duplicada": duplicados,
    }
    catalogo = list(indice_rpc.values())
    rutas_existentes = rutas_existentes or {}
    total_items = len(items)
    log.info(
        "Reconciliación global: %d metadata(s); reorganizar_output=%s",
        total_items,
        reorganizar_output,
    )

    for indice_item, item in enumerate(items, start=1):
        if indice_item == 1 or indice_item % 100 == 0 or indice_item == total_items:
            log.info(
                "Reconciliación global: procesando metadata %d/%d",
                indice_item,
                total_items,
            )
        carpeta: Path = item["carpeta"]
        meta_satys = item["meta_satys"]
        meta_tn = item["meta_tn"]
        registro = item["registro"]
        folio = _folio(meta_satys, meta_tn, registro)
        folio_opc = folio_opc_desde_metadata(meta_satys, meta_tn)
        identificador = _identificador_salida(carpeta, descargas_base, registro)
        id_solicitante = bc.normalizar_id(
            meta_satys.get("id_solicitante") or meta_tn.get("id_solicitante")
        )
        nombre_satys = (
            meta_satys.get("nombre_operador")
            or meta_satys.get("concesionario")
            or meta_tn.get("nombre_operador")
            or meta_tn.get("concesionario")
            or ""
        )
        ruta_existente = normalizar_ruta_excel(rutas_existentes.get(registro, ""))
        if ruta_es_canonica_estable(ruta_existente):
            # La reconciliación diaria no vuelve a resolver ni toca registros
            # históricos que ya tienen una Ruta canónica validada.
            resolucion = {
                "ok": True,
                "metodo": "ruta_excel_preservada",
                "ruta": ruta_existente,
                "nombre_completo": nombre_satys,
                "idBp": id_solicitante,
                "score": 1.0,
            }
        else:
            resolucion = bc.resolver_operador_seguro(
                id_solicitante,
                nombre_satys,
                catalogo,
            )

        resultado: dict[str, Any] = {
            "folio": folio,
            "folio_id": identificador,
            "folio_opc": folio_opc,
            "registro": registro,
            "descargas_dir": str(carpeta),
            "id_solicitante": id_solicitante,
            "nombre_operador": nombre_satys,
            "representante_legal": (
                meta_satys.get("representante_legal")
                or meta_tn.get("representante_legal")
                or ""
            ),
        }

        if es_folio_opc_correo(folio_opc):
            identificador_correo = str(registro or identificador)
            ruta_operador = ""
            if resolucion.get("ok"):
                if resolucion.get("operadores"):
                    ruta_operador = construir_ruta_operadores(
                        resolucion["operadores"],
                        registro or identificador,
                    )
                else:
                    ruta_operador = construir_ruta(
                        resolucion["nombre_completo"],
                        resolucion["idBp"],
                        registro or identificador,
                    )
            resolucion["ruta_operador_rpc"] = ruta_operador
            resolucion["ruta"] = ruta_relativa_sin_operador(
                identificador_correo,
                folio_opc,
            )
            destino = destino_sin_operador(
                output_base,
                identificador_correo,
                folio_opc,
            )
            organizacion = (
                organizar_correo_exclusivo(
                    carpeta,
                    output_base,
                    identificador_correo,
                    folio_opc=folio_opc,
                    ruta_operador=ruta_operador,
                    identificadores_legacy=(identificador,),
                )
                if migrar_correos and reorganizar_output
                else {
                    "destino": destino,
                    "archivos_copiados": [],
                    "verificado": destino.is_dir(),
                    "duplicados_retirados": [],
                    "errores": [],
                }
            )
            resultado.update({
                "es_correo": True,
                "identificador_correo": identificador_correo,
                "rpc_ok": bool(resolucion.get("ok")),
                "organizado_ok": organizacion["verificado"],
                "archivos_copiados": len(organizacion["archivos_copiados"]),
                "nombre_operador": (
                    resolucion.get("nombre_completo") or nombre_satys
                ),
                "rpc_resultado": resolucion,
                "sin_operador_dir": str(organizacion["destino"]),
                "output_dir": str(organizacion["destino"]),
                "duplicados_correo_retirados": organizacion["duplicados_retirados"],
                "errores_organizacion_correo": organizacion["errores"],
            })
            stats["sin_operador_correo"] += 1
            stats["archivos_correo_copiados"] += len(
                organizacion["archivos_copiados"]
            )
        elif resolucion.get("ok"):
            if resolucion.get("metodo") == "ruta_excel_preservada":
                ruta = ruta_existente
            elif resolucion.get("operadores"):
                ruta = construir_ruta_operadores(
                    resolucion["operadores"],
                    registro or identificador,
                )
            else:
                ruta = construir_ruta(
                    resolucion["nombre_completo"],
                    resolucion["idBp"],
                    registro or identificador,
                )
            destino = output_base / ruta.replace("\\", "/")
            if reorganizar_output:
                destino_organizado = organizar_archivos(
                    carpeta,
                    ruta,
                    output_base=output_base,
                )
                copiados = (
                    sum(1 for _ in iter_archivos_publicables_output(carpeta))
                    if destino_organizado is not None
                    else 0
                )
                organizado_ok = destino_organizado is not None
            else:
                # El pipeline principal ya organizó descargas -> output. La
                # reconciliación diaria sólo necesita reconstruir Excel/Ruta;
                # volver a copiar y comparar byte a byte decenas de miles de
                # archivos históricos puede tardar horas y no aporta datos.
                copiados = 0
                organizado_ok = destino.is_dir()
            resolucion["ruta"] = ruta
            resultado.update({
                "rpc_ok": True,
                "organizado_ok": organizado_ok,
                "archivos_copiados": copiados,
                "nombre_operador": resolucion["nombre_completo"],
                "rpc_resultado": resolucion,
                "output_dir": str(destino),
            })
            stats["rpc_ok"] += 1
            stats["archivos_organizados"] += copiados
        else:
            destino = destino_sin_operador(output_base, identificador, folio_opc)
            organizacion = (
                organizar_correo_exclusivo(
                    carpeta,
                    output_base,
                    identificador,
                    folio_opc=folio_opc,
                    identificadores_legacy=(registro,),
                )
                if migrar_correos and reorganizar_output
                else {
                    "destino": destino,
                    "archivos_copiados": [],
                    "verificado": destino.is_dir(),
                    "duplicados_retirados": [],
                    "errores": [],
                }
            )
            resultado.update({
                "rpc_ok": False,
                "organizado_ok": organizacion["verificado"],
                "archivos_copiados": len(organizacion["archivos_copiados"]),
                "rpc_resultado": resolucion,
                "sin_operador_dir": str(organizacion["destino"]),
                "output_dir": str(organizacion["destino"]),
                "duplicados_revision_retirados": organizacion["duplicados_retirados"],
                "errores_organizacion_revision": organizacion["errores"],
            })
            stats["sin_operador"] += 1

        resultados.append(resultado)

    return resultados, stats


def ejecutar(
    *,
    descargas_base: Path,
    output_base: Path,
    excel_path: Path,
    base_rpc: Path,
    project_root: Path,
    migrar_correos: bool = True,
    reorganizar_output: bool = True,
    crear_backup: bool = True,
) -> dict[str, Any]:
    indice_rpc, excel_rpc = cargar_indice_rpc(base_rpc)
    rutas_existentes = cargar_rutas_maestro(excel_path)
    resultados, stats = construir_resultados(
        descargas_base,
        output_base,
        indice_rpc,
        rutas_existentes=rutas_existentes,
        migrar_correos=migrar_correos,
        reorganizar_output=reorganizar_output,
    )
    if not resultados:
        raise RuntimeError(f"No se encontraron metadatos válidos bajo {descargas_base}.")

    from reporte_operadores import generar_reportes_operadores

    reportes_csv = generar_reportes_operadores(
        resultados,
        modo="reconciliacion_global",
        logs_dir=project_root / "logs",
    )

    consolidado = generar_excel_metadata_json(
        resultados=resultados,
        descargas_base=descargas_base,
        output_base=output_base,
        excel_salida=output_base / "Folios_Datos_Completos.xlsx",
        project_root=project_root,
    )
    reconciliacion = aplicar_rutas_conservadoras(
        excel_path, resultados, crear_backup=crear_backup
    )
    log.info(
        "Reconciliación conservadora: %d Ruta(s) promovidas; %d canónicas preservadas; %d pendientes.",
        reconciliacion["routes_promoted"],
        reconciliacion["canonical_preserved"],
        reconciliacion["still_pending"],
    )

    return {
        "ok": True,
        "fecha": datetime.now().isoformat(),
        "excel_rpc": str(excel_rpc) if excel_rpc else "",
        "excel_consolidado": str(consolidado),
        "excel_maestro": str(excel_path),
        "estadisticas": stats,
        "reportes_csv": reportes_csv,
        "reconciliacion": reconciliacion,
    }


def construir_parser() -> argparse.ArgumentParser:
    project_root = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Promueve de forma conservadora Rutas pendientes sin sobrescribir el histórico validado."
    )
    parser.add_argument("--descargas", type=Path, default=ruta_configurada("descargas", "descargas"))
    parser.add_argument("--output", type=Path, default=ruta_configurada("output", "output"))
    parser.add_argument("--excel", type=Path, default=ruta_configurada("excel", "TrámitesCRT.xlsx"))
    parser.add_argument("--base-rpc", type=Path, default=project_root / "base_de_datos_rpc")
    parser.add_argument("--resumen-json", type=Path, default=project_root / "logs" / "reconciliacion_global_ultimo.json")
    parser.add_argument("--sin-migrar-correos", action="store_true")
    parser.add_argument(
        "--sin-reorganizar-output",
        action="store_true",
        help=(
            "No vuelve a copiar/verificar todos los archivos históricos de descargas a output; "
            "sólo reconstruye rutas, reportes y Excel. Recomendado para la corrida diaria."
        ),
    )
    parser.add_argument("--sin-backup", action="store_true")
    return parser


def main() -> int:
    args = construir_parser().parse_args()
    args.resumen_json.parent.mkdir(parents=True, exist_ok=True)
    try:
        resumen = ejecutar(
            descargas_base=args.descargas,
            output_base=args.output,
            excel_path=args.excel,
            base_rpc=args.base_rpc,
            project_root=Path(__file__).resolve().parent,
            migrar_correos=not args.sin_migrar_correos,
            reorganizar_output=not args.sin_reorganizar_output,
            crear_backup=not args.sin_backup,
        )
        args.resumen_json.write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(resumen, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        resumen = {"ok": False, "fecha": datetime.now().isoformat(), "error": str(exc)}
        args.resumen_json.write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")
        log.exception("Reconciliación global fallida: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
