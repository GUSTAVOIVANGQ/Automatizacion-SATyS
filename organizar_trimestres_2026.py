#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Organización adicional de expedientes SATyS por ventanas Q3/Q4 solicitadas.

Reglas de negocio (mapeo deliberadamente personalizado):
  - 2026-10-01 .. 2026-12-15 -> output/2026Q3/<Ruta base>
  - 2027-01-01 .. 2027-03-31 -> output/2026Q4/<Ruta base>

La fuente de verdad de los archivos siempre es ``descargas``. La columna ``1711``
identifica el expediente; la columna ``Ruta`` aporta la jerarquía canónica y
``Fecha de creación`` decide el bucket. Para los registros dentro de las ventanas,
la etapa garantiza DOS copias organizadas: ``output/<Ruta base>`` y
``output/<bucket>/<Ruta base>``. Después actualiza ``Ruta`` en TrámitesCRT.xlsx a
``<bucket>\\<Ruta base>`` (Ruta es relativa a ``output``) y replica output + Excel
a DEPI antes del correo.

La operación es no destructiva e idempotente:
  * nunca mueve ni borra ``descargas``;
  * nunca inventa sufijos ``_1``/``_2`` para colisiones;
  * un archivo actual de ``descargas`` sobrescribe el mismo path en ambos destinos;
  * archivos distintos ya existentes se conservan;
  * JSON permanece sólo en ``descargas`` y no se publica en ``output``;
  * un prefijo 2026Q3/2026Q4 ya existente en Ruta se retira antes de reconstruirla;
  * con DEPI habilitado, Ruta sólo se confirma cuando sus dos destinos se publicaron.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import openpyxl

from configuracion_local import carpeta_compartida, ruta_configurada
from proceso_lock import LockOcupadoError, ProcesoLock
from sincronizacion_depi import validar_destino_compartido
from guardado_seguro import reemplazar_desde_temporal

PROJECT_DIR = Path(__file__).resolve().parent
LOGS_DIR_DEFAULT = PROJECT_DIR / "logs"
HOJAS_DEFAULT = ("Turnados recibidos", "Internos")

# IMPORTANTE: este no es el trimestre calendario estándar. Es el mapeo exacto
# solicitado para la exportación adicional.
VENTANAS_2026 = (
    ("2026Q3", date(2026, 10, 1), date(2026, 12, 15)),
    ("2026Q4", date(2027, 1, 1), date(2027, 3, 31)),
)


@dataclass
class Resultado:
    filas_excel: int = 0
    seleccionados: dict[str, int] = field(default_factory=lambda: {"2026Q3": 0, "2026Q4": 0})
    expedientes_con_fuente: int = 0
    expedientes_sin_fuente: int = 0
    rutas_invalidas: int = 0
    archivos_locales_copiados: int = 0
    archivos_locales_iguales: int = 0
    archivos_normales_copiados: int = 0
    archivos_normales_iguales: int = 0
    archivos_bucket_copiados: int = 0
    archivos_bucket_iguales: int = 0
    archivos_depi_copiados: int = 0
    archivos_depi_iguales: int = 0
    rutas_excel_actualizadas: int = 0
    rutas_excel_ya_correctas: int = 0
    rutas_excel_bloqueadas_depi: int = 0
    excel_depi_actualizado: bool = False
    excel_backup: str = ""
    json_omitidos: int = 0
    errores: list[str] = field(default_factory=list)
    detalle: list[dict] = field(default_factory=list)


def _normalizar_cabecera(valor) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto).strip().casefold()
    return texto


def _indices_columnas_desde_valores(encabezados, titulo_hoja: str) -> dict[str, int]:
    """Devuelve índices 0-based de las columnas requeridas.

    IMPORTANTE: en un ``ReadOnlyWorksheet`` no se debe usar ``ws.cell(row, col)``
    repetidamente. OpenPyXL vuelve a abrir y recorrer el XML de la hoja para cada
    consulta aleatoria, lo que convierte un Excel de pocos miles de filas en una
    operación de complejidad cuadrática y puede tardar horas. Esta función recibe
    el encabezado ya obtenido por ``iter_rows(values_only=True)`` para mantener una
    sola pasada secuencial por la hoja.
    """
    buscadas = {
        "1711": "1711",
        "ruta": "ruta",
        "fecha": "fecha de creacion",
    }
    halladas: dict[str, int] = {}
    for idx, valor in enumerate(encabezados):
        norm = _normalizar_cabecera(valor)
        for clave, esperado in buscadas.items():
            if norm == esperado:
                halladas[clave] = idx
    faltan = sorted(set(buscadas) - set(halladas))
    if faltan:
        raise ValueError(f"Hoja {titulo_hoja!r}: faltan columnas requeridas: {', '.join(faltan)}")
    return halladas


def _fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto:
        return None
    formatos = (
        "%d/%m/%Y",
        "%d/%m/%Y %H:%M:%S",
        "%d-%m-%Y",
        "%d-%m-%Y %H:%M:%S",
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
    )
    for fmt in formatos:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def bucket_para_fecha(fecha: date | None) -> str | None:
    if fecha is None:
        return None
    for nombre, inicio, fin in VENTANAS_2026:
        if inicio <= fecha <= fin:
            return nombre
    return None


def _ruta_relativa_segura(valor) -> Path:
    if valor is None:
        raise ValueError("Ruta vacía")
    texto = str(valor).strip().replace("\\", "/")
    texto = re.sub(r"/+", "/", texto)
    # Tolera rutas absolutas antiguas mostradas en reportes, pero las vuelve
    # relativas a output antes de reproducirlas bajo 2026Q3/2026Q4.
    plegada = texto.casefold()
    for prefijo in ("/app/output/", "app/output/", "output/"):
        if plegada.startswith(prefijo):
            texto = texto[len(prefijo):]
            break
    texto = texto.strip("/ ")
    if not texto:
        raise ValueError("Ruta vacía después de normalizar")
    partes = [p.strip() for p in texto.split("/") if p.strip()]
    if not partes or any(p in {".", ".."} for p in partes):
        raise ValueError(f"Ruta insegura: {valor!r}")
    if partes[0].casefold() in {"2026q3", "2026q4"}:
        # Evita crear output/2026Qx/2026Qx/... si el Excel llegara a contener
        # por accidente una ruta de exportación en vez de la Ruta maestra.
        partes = partes[1:]
    if not partes:
        raise ValueError(f"Ruta sin destino canónico: {valor!r}")
    return Path(*partes)


def _indice_descargas(descargas: Path) -> dict[str, list[Path]]:
    indice: dict[str, list[Path]] = {}
    for raiz, dirs, _files in os.walk(descargas):
        # Evita directorios técnicos que no son expedientes.
        dirs[:] = [d for d in dirs if d not in {"__pycache__", ".git"}]
        for nombre in dirs:
            p = Path(raiz) / nombre
            indice.setdefault(nombre.casefold(), []).append(p)
    for lista in indice.values():
        lista.sort(key=lambda p: (p.stat().st_mtime_ns if p.exists() else 0, str(p)))
    return indice


def _fuentes_para(indice: dict[str, list[Path]], identificador: str) -> list[Path]:
    key = identificador.strip().casefold()
    exactas = list(indice.get(key, []))
    if exactas:
        return _sin_ancestros_repetidos(exactas)
    # Internos puede tener un sufijo técnico (p. ej. 118349_9d109380).
    prefijo = key + "_"
    candidatas: list[Path] = []
    for nombre, paths in indice.items():
        if nombre.startswith(prefijo):
            candidatas.extend(paths)
    return _sin_ancestros_repetidos(candidatas)


def _sin_ancestros_repetidos(paths: Iterable[Path]) -> list[Path]:
    # Primero elimina carpetas hijas redundantes; después devuelve las fuentes
    # por antigüedad para que la más reciente se copie al final y gane colisiones.
    unicos = sorted({p.resolve() for p in paths}, key=lambda p: (len(p.parts), str(p)))
    salida: list[Path] = []
    for p in unicos:
        if any(q == p or q in p.parents for q in salida):
            continue
        salida.append(p)
    salida.sort(key=lambda p: (p.stat().st_mtime_ns, str(p)))
    return salida


def _archivos_fuente(fuente: Path) -> Iterable[tuple[Path, Path]]:
    for item in sorted(fuente.rglob("*"), key=lambda p: str(p).casefold()):
        if item.is_symlink() or not item.is_file():
            continue
        relativo = item.relative_to(fuente)
        if item.suffix.casefold() == ".json":
            yield item, Path("__OMITIR_JSON__") / relativo
            continue
        yield item, relativo


def _misma_version(src: Path, dst: Path) -> bool:
    """Compara contenido real, no sólo tamaño/mtime.

    ``descargas`` es la fuente de verdad. Dos archivos distintos pueden tener el
    mismo tamaño y el mismo mtime (por ejemplo, tras restauraciones/copy2). Si se
    trataban como iguales por metadatos, el bucket podía conservar una versión
    vieja. La comparación por bloques garantiza que una diferencia de contenido
    provoque la sobrescritura del pathname canónico, sin crear sufijos.
    """
    try:
        if not dst.is_file():
            return False
        if src.stat().st_size != dst.stat().st_size:
            return False
        with src.open('rb') as a, dst.open('rb') as b:
            while True:
                ca = a.read(1024 * 1024)
                cb = b.read(1024 * 1024)
                if ca != cb:
                    return False
                if not ca:
                    return True
    except OSError:
        return False


def _destino_bajo(root: Path, rel: Path) -> Path:
    destino = root / rel
    raiz_resuelta = root.resolve()
    padre_resuelto = destino.parent.resolve(strict=False)
    if padre_resuelto != raiz_resuelta and raiz_resuelta not in padre_resuelto.parents:
        raise ValueError(f"Destino fuera de raíz permitida: {destino}")
    return destino


def _copiar_si_necesario(src: Path, dst: Path) -> bool:
    """Copia src a dst. Devuelve True si escribió, False si ya era idéntico."""
    if _misma_version(src, dst):
        return False
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    return True


def _ruta_excel_bucket(bucket: str, ruta_rel: Path) -> str:
    """Ruta que se guarda en Excel; siempre relativa a output e idempotente."""
    return str(Path(bucket) / ruta_rel).replace("/", "\\")


def _copiar_arbol_fuentes(
    *,
    fuentes: list[Path],
    destino_normal: Path,
    destino_bucket: Path,
    resultado: Resultado,
    detalle: dict,
) -> None:
    """Publica una fuente en los DOS destinos locales, con descargas como verdad."""
    destino_normal.mkdir(parents=True, exist_ok=True)
    destino_bucket.mkdir(parents=True, exist_ok=True)
    for fuente in fuentes:
        for src, relativo in _archivos_fuente(fuente):
            if relativo.parts and relativo.parts[0] == "__OMITIR_JSON__":
                resultado.json_omitidos += 1
                continue

            dst_normal = _destino_bajo(destino_normal, relativo)
            if _copiar_si_necesario(src, dst_normal):
                resultado.archivos_normales_copiados += 1
                resultado.archivos_locales_copiados += 1
            else:
                resultado.archivos_normales_iguales += 1
                resultado.archivos_locales_iguales += 1

            dst_bucket = _destino_bajo(destino_bucket, relativo)
            if _copiar_si_necesario(src, dst_bucket):
                resultado.archivos_bucket_copiados += 1
                resultado.archivos_locales_copiados += 1
            else:
                resultado.archivos_bucket_iguales += 1
                resultado.archivos_locales_iguales += 1
            detalle["archivos"] += 1


def _sincronizar_directorio(local_dir: Path, remote_dir: Path, resultado: Resultado) -> None:
    """Merge completo local->DEPI, excluyendo JSON y sobrescribiendo mismo pathname."""
    for src in sorted(local_dir.rglob("*"), key=lambda p: str(p).casefold()):
        if src.is_symlink() or not src.is_file():
            continue
        if src.suffix.casefold() == ".json":
            resultado.json_omitidos += 1
            continue
        rel = src.relative_to(local_dir)
        dst = _destino_bajo(remote_dir, rel)
        if _copiar_si_necesario(src, dst):
            resultado.archivos_depi_copiados += 1
        else:
            resultado.archivos_depi_iguales += 1


def _actualizar_excel_rutas(
    excel_path: Path,
    actualizaciones: list[dict],
    resultado: Resultado,
) -> Path | None:
    """Actualiza Ruta con backup + sustitución segura del Excel bind-mounted."""
    if not actualizaciones:
        return None
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = excel_path.with_name(
        f"{excel_path.stem}_backup_pre_trimestres_{sello}{excel_path.suffix}"
    )
    shutil.copy2(excel_path, backup)
    resultado.excel_backup = str(backup)

    wb = openpyxl.load_workbook(excel_path, read_only=False, data_only=False)
    cambios = 0
    ya_correctas = 0
    try:
        for upd in actualizaciones:
            ws = wb[upd["hoja"]]
            cell = ws.cell(row=upd["fila"], column=upd["col_ruta"] + 1)
            actual = "" if cell.value is None else str(cell.value).strip()
            if actual == upd["ruta_nueva"]:
                ya_correctas += 1
            else:
                cell.value = upd["ruta_nueva"]
                cambios += 1
        temporal = excel_path.with_name(
            f".{excel_path.stem}.trimestres_{sello}.tmp{excel_path.suffix}"
        )
        wb.save(temporal)
    finally:
        wb.close()

    reemplazar_desde_temporal(temporal, excel_path)
    resultado.rutas_excel_actualizadas = cambios
    resultado.rutas_excel_ya_correctas = ya_correctas
    return backup


def organizar(
    *,
    excel_path: Path,
    descargas_base: Path,
    output_base: Path,
    shared_root: Path,
    logs_dir: Path,
    dry_run: bool = False,
    sincronizar_depi: bool = True,
) -> dict:
    excel_path = Path(excel_path)
    descargas_base = Path(descargas_base)
    output_base = Path(output_base)
    shared_root = Path(shared_root)
    logs_dir = Path(logs_dir)

    if not excel_path.is_file():
        raise FileNotFoundError(f"No existe Excel: {excel_path}")
    if not descargas_base.is_dir():
        raise FileNotFoundError(f"No existe descargas: {descargas_base}")
    if not dry_run:
        output_base.mkdir(parents=True, exist_ok=True)
    logs_dir.mkdir(parents=True, exist_ok=True)

    resultado = Resultado()
    sync_solicitado = bool(sincronizar_depi and not dry_run)
    sync_disponible = False
    if sync_solicitado:
        error_shared = validar_destino_compartido(shared_root)
        if error_shared:
            resultado.errores.append(error_shared)
        else:
            (shared_root / "output").mkdir(parents=True, exist_ok=True)
            sync_disponible = True

    print("📚 Indexando carpetas de descargas una sola vez...")
    indice = _indice_descargas(descargas_base)
    print(f"📚 Índice listo: {sum(len(v) for v in indice.values())} carpeta(s), {len(indice)} nombre(s) distinto(s).")

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=False)
    seleccionados: list[dict] = []
    try:
        for hoja in HOJAS_DEFAULT:
            if hoja not in wb.sheetnames:
                continue
            ws = wb[hoja]
            filas_iter = ws.iter_rows(values_only=True)
            try:
                encabezados = next(filas_iter)
            except StopIteration:
                continue
            cols = _indices_columnas_desde_valores(encabezados, hoja)
            filas_hoja = 0
            for fila, valores in enumerate(filas_iter, start=2):
                filas_hoja += 1
                resultado.filas_excel += 1
                ident = valores[cols["1711"]] if cols["1711"] < len(valores) else None
                fecha_val = valores[cols["fecha"]] if cols["fecha"] < len(valores) else None
                ruta_val = valores[cols["ruta"]] if cols["ruta"] < len(valores) else None
                if ident is not None:
                    fecha = _fecha(fecha_val)
                    bucket = bucket_para_fecha(fecha)
                    if bucket:
                        resultado.seleccionados[bucket] += 1
                        seleccionados.append({
                            "hoja": hoja,
                            "fila": fila,
                            "1711": str(ident).strip(),
                            "fecha": fecha,
                            "bucket": bucket,
                            "ruta_original": "" if ruta_val is None else str(ruta_val),
                            "col_ruta": cols["ruta"],
                        })
                if filas_hoja % 500 == 0:
                    print(
                        f"📊 Excel {hoja}: {filas_hoja} fila(s) leídas; "
                        f"seleccionados acumulados={len(seleccionados)}."
                    )
            print(
                f"📊 Excel {hoja} terminado: {filas_hoja} fila(s) leídas; "
                f"seleccionados acumulados={len(seleccionados)}."
            )
    finally:
        wb.close()

    seleccionados.sort(key=lambda x: (x["fecha"], x["1711"], x["hoja"], x["fila"]))
    rutas_tocadas: set[tuple[str, str]] = set()
    candidatos_excel: list[dict] = []

    for n, item in enumerate(seleccionados, start=1):
        ident = item["1711"]
        bucket = item["bucket"]
        detalle = {
            "hoja": item["hoja"], "fila": item["fila"], "1711": ident,
            "fecha": item["fecha"].isoformat(), "bucket": bucket,
            "ruta_excel": item["ruta_original"], "ruta_base": "", "ruta_nueva": "",
            "fuentes": [], "estado": "pendiente", "archivos": 0,
        }
        try:
            ruta_rel = _ruta_relativa_segura(item["ruta_original"])
            ruta_nueva = _ruta_excel_bucket(bucket, ruta_rel)
            detalle["ruta_base"] = str(ruta_rel).replace("/", "\\")
            detalle["ruta_nueva"] = ruta_nueva
        except Exception as exc:
            resultado.rutas_invalidas += 1
            detalle["estado"] = "ruta_invalida"
            detalle["error"] = str(exc)
            resultado.detalle.append(detalle)
            print(f"⚠️ [{n}/{len(seleccionados)}] {ident}: Ruta inválida: {exc}")
            continue

        fuentes = _fuentes_para(indice, ident)
        detalle["fuentes"] = [str(p) for p in fuentes]
        if not fuentes:
            resultado.expedientes_sin_fuente += 1
            detalle["estado"] = "sin_fuente_descargas"
            resultado.detalle.append(detalle)
            print(f"⚠️ [{n}/{len(seleccionados)}] {ident}: no se encontró carpeta fuente en descargas.")
            continue
        resultado.expedientes_con_fuente += 1

        destino_normal = output_base / ruta_rel
        destino_bucket = output_base / bucket / ruta_rel
        detalle["destino_normal"] = str(destino_normal)
        detalle["destino_bucket"] = str(destino_bucket)
        if dry_run:
            detalle["estado"] = "dry_run"
            resultado.detalle.append(detalle)
            print(
                f"🔎 [{n}/{len(seleccionados)}] {ident} -> output/{ruta_rel} + "
                f"output/{bucket}/{ruta_rel}; Ruta => {ruta_nueva}"
            )
            continue

        try:
            _copiar_arbol_fuentes(
                fuentes=fuentes,
                destino_normal=destino_normal,
                destino_bucket=destino_bucket,
                resultado=resultado,
                detalle=detalle,
            )
            rutas_tocadas.add(("normal", ruta_rel.as_posix()))
            rutas_tocadas.add((bucket, ruta_rel.as_posix()))
            detalle["estado"] = "organizado_local"
            candidatos_excel.append({
                "hoja": item["hoja"], "fila": item["fila"], "col_ruta": item["col_ruta"],
                "ruta_nueva": ruta_nueva, "bucket": bucket, "ruta_rel": ruta_rel.as_posix(),
                "detalle": detalle,
            })
            resultado.detalle.append(detalle)
            print(
                f"✅ [{n}/{len(seleccionados)}] {ident} -> output/{ruta_rel} + "
                f"output/{bucket}/{ruta_rel} ({detalle['archivos']} archivo(s)); Ruta => {ruta_nueva}"
            )
        except Exception as exc:
            msg = f"{ident} -> {bucket}/{ruta_rel}: {type(exc).__name__}: {exc}"
            resultado.errores.append(msg)
            detalle["estado"] = "error"
            detalle["error"] = str(exc)
            resultado.detalle.append(detalle)
            print(f"❌ [{n}/{len(seleccionados)}] {msg}", file=sys.stderr)

    sync_fallos: set[tuple[str, str]] = set()
    if sync_solicitado and sync_disponible:
        for tipo, ruta_texto in sorted(rutas_tocadas):
            ruta_rel = Path(ruta_texto)
            if tipo == "normal":
                local_dir = output_base / ruta_rel
                remote_dir = shared_root / "output" / ruta_rel
            else:
                local_dir = output_base / tipo / ruta_rel
                remote_dir = shared_root / "output" / tipo / ruta_rel
            try:
                _sincronizar_directorio(local_dir, remote_dir, resultado)
            except Exception as exc:
                key = (tipo, ruta_texto)
                sync_fallos.add(key)
                msg = f"DEPI {tipo}/{ruta_rel}: {type(exc).__name__}: {exc}"
                resultado.errores.append(msg)
                print(f"❌ {msg}", file=sys.stderr)

    actualizaciones: list[dict] = []
    for cand in candidatos_excel:
        requiere_sync_ok = sync_solicitado
        normal_key = ("normal", cand["ruta_rel"])
        bucket_key = (cand["bucket"], cand["ruta_rel"])
        if requiere_sync_ok and (not sync_disponible or normal_key in sync_fallos or bucket_key in sync_fallos):
            resultado.rutas_excel_bloqueadas_depi += 1
            cand["detalle"]["estado"] = "organizado_local_sin_confirmar_depi"
            continue
        cand["detalle"]["estado"] = "listo_para_excel"
        actualizaciones.append(cand)

    backup = None
    if actualizaciones and not dry_run:
        backup = _actualizar_excel_rutas(excel_path, actualizaciones, resultado)
        for cand in actualizaciones:
            cand["detalle"]["estado"] = "ruta_excel_actualizada"

        if sync_solicitado and sync_disponible:
            try:
                # El Excel compartido debe reflejar exactamente las Rutas ya confirmadas.
                dst_excel = shared_root / excel_path.name
                if _copiar_si_necesario(excel_path, dst_excel):
                    resultado.excel_depi_actualizado = True
                else:
                    resultado.excel_depi_actualizado = True
            except Exception as exc:
                msg = f"DEPI Excel {excel_path.name}: {type(exc).__name__}: {exc}"
                resultado.errores.append(msg)
                print(f"❌ {msg}", file=sys.stderr)
                # Evita dejar el Excel local anunciando una Ruta trimestral que no
                # pudo publicarse en el Excel de DEPI. Los archivos ya copiados se
                # conservan y la siguiente ejecución puede reintentar idempotentemente.
                if backup and backup.is_file():
                    shutil.copyfile(backup, excel_path)
                    resultado.rutas_excel_actualizadas = 0
                    resultado.rutas_excel_ya_correctas = 0
                    resultado.excel_depi_actualizado = False
                    for cand in actualizaciones:
                        cand["detalle"]["estado"] = "excel_rollback_por_sync_depi"

    payload = {
        "fecha": datetime.now().isoformat(),
        "excel": str(excel_path),
        "descargas": str(descargas_base),
        "output": str(output_base),
        "shared": str(shared_root),
        "dry_run": dry_run,
        "sincronizar_depi": bool(sincronizar_depi),
        "mapeo": {
            "2026Q3": {"desde": "2026-10-01", "hasta": "2026-12-15"},
            "2026Q4": {"desde": "2027-01-01", "hasta": "2027-03-31"},
        },
        "regla_ruta_excel": "Ruta=<bucket>\\<Ruta base>; Ruta es relativa a output",
        "filas_excel": resultado.filas_excel,
        "seleccionados": resultado.seleccionados,
        "total_seleccionados": sum(resultado.seleccionados.values()),
        "expedientes_con_fuente": resultado.expedientes_con_fuente,
        "expedientes_sin_fuente": resultado.expedientes_sin_fuente,
        "rutas_invalidas": resultado.rutas_invalidas,
        "archivos_locales_copiados": resultado.archivos_locales_copiados,
        "archivos_locales_iguales": resultado.archivos_locales_iguales,
        "archivos_normales_copiados": resultado.archivos_normales_copiados,
        "archivos_normales_iguales": resultado.archivos_normales_iguales,
        "archivos_bucket_copiados": resultado.archivos_bucket_copiados,
        "archivos_bucket_iguales": resultado.archivos_bucket_iguales,
        "archivos_depi_copiados": resultado.archivos_depi_copiados,
        "archivos_depi_iguales": resultado.archivos_depi_iguales,
        "rutas_excel_actualizadas": resultado.rutas_excel_actualizadas,
        "rutas_excel_ya_correctas": resultado.rutas_excel_ya_correctas,
        "rutas_excel_bloqueadas_depi": resultado.rutas_excel_bloqueadas_depi,
        "excel_backup": resultado.excel_backup,
        "excel_depi_actualizado": resultado.excel_depi_actualizado,
        "json_omitidos": resultado.json_omitidos,
        "errores": resultado.errores,
        "detalle": resultado.detalle,
    }

    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = logs_dir / f"organizar_trimestres_2026_{sello}.json"
    latest_path = logs_dir / "organizar_trimestres_2026_ultimo.json"
    csv_path = logs_dir / f"organizar_trimestres_2026_{sello}.csv"
    texto = json.dumps(payload, ensure_ascii=False, indent=2)
    json_path.write_text(texto, encoding="utf-8")
    latest_path.write_text(texto, encoding="utf-8")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "hoja", "fila", "1711", "fecha", "bucket", "ruta_excel", "ruta_base",
                "ruta_nueva", "estado", "archivos", "destino_normal", "destino_bucket", "error",
            ],
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(resultado.detalle)

    print("\n" + "=" * 78)
    print("ORGANIZACIÓN ADICIONAL Q3/Q4 TERMINADA")
    print(f"2026Q3 (01-oct-2026..15-dic-2026): {resultado.seleccionados['2026Q3']} registro(s)")
    print(f"2026Q4 (01-ene-2027..31-mar-2027): {resultado.seleccionados['2026Q4']} registro(s)")
    print(f"Fuentes encontradas: {resultado.expedientes_con_fuente}; sin fuente: {resultado.expedientes_sin_fuente}")
    print(
        f"Normal output: escritos={resultado.archivos_normales_copiados}, iguales={resultado.archivos_normales_iguales}; "
        f"bucket: escritos={resultado.archivos_bucket_copiados}, iguales={resultado.archivos_bucket_iguales}"
    )
    if sync_solicitado:
        print(f"Archivos DEPI escritos: {resultado.archivos_depi_copiados}; ya iguales: {resultado.archivos_depi_iguales}")
    print(
        f"Rutas Excel actualizadas: {resultado.rutas_excel_actualizadas}; ya correctas: {resultado.rutas_excel_ya_correctas}; "
        f"bloqueadas por DEPI: {resultado.rutas_excel_bloqueadas_depi}"
    )
    print(f"JSON omitidos de output: {resultado.json_omitidos}; errores: {len(resultado.errores)}")
    print(f"Resumen: {json_path}")
    print("=" * 78)
    return payload

def construir_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description=(
            "Mantiene expedientes en output/<Ruta base> y además en output/2026Q3 o output/2026Q4 "
            "según Fecha de creación; actualiza Ruta en TrámitesCRT.xlsx y sincroniza output + Excel a DEPI."
        )
    )
    p.add_argument("--excel", type=Path, default=ruta_configurada("excel", "TrámitesCRT.xlsx"))
    p.add_argument("--descargas", type=Path, default=ruta_configurada("descargas", "descargas"))
    p.add_argument("--output", type=Path, default=ruta_configurada("output", "output"))
    p.add_argument("--shared", type=Path, default=carpeta_compartida())
    p.add_argument("--logs-dir", type=Path, default=LOGS_DIR_DEFAULT)
    p.add_argument("--dry-run", action="store_true", help="Sólo calcula; no escribe output, Excel ni DEPI.")
    p.add_argument("--sin-sync-depi", action="store_true", help="Organiza localmente pero no replica los buckets a DEPI.")
    p.add_argument("--sin-lock", action="store_true", help="Sólo para subproceso/pruebas; no usar manualmente en producción.")
    return p


def main() -> int:
    args = construir_parser().parse_args()
    lock = None
    try:
        if not args.sin_lock:
            lock = ProcesoLock(proceso="organizar_trimestres_2026.py")
            lock.adquirir()
            print("🔒 Lock global SATyS adquirido para organización adicional Q3/Q4.")
        payload = organizar(
            excel_path=args.excel,
            descargas_base=args.descargas,
            output_base=args.output,
            shared_root=args.shared,
            logs_dir=args.logs_dir,
            dry_run=args.dry_run,
            sincronizar_depi=not args.sin_sync_depi,
        )
        # Omisiones de fuente/ruta se auditan pero no convierten la corrida diaria
        # en fallo. Sólo errores de E/S/DEPI reales devuelven código 2.
        return 0 if not payload["errores"] else 2
    except LockOcupadoError as exc:
        print(f"ERROR: SATyS está ocupado: {exc}", file=sys.stderr)
        return 3
    except Exception as exc:
        print(f"ERROR fatal en organización adicional Q3/Q4: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    finally:
        if lock is not None:
            lock.liberar()


if __name__ == "__main__":
    raise SystemExit(main())
